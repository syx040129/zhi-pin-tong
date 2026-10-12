from pathlib import Path
import re, json, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(r'G:\project\CareerNova\thesis-workspace')
source = root / 'workspace/drafts/开题报告.md'
out = root / 'workspace/final/智聘通招聘网站的设计与实现_开题报告.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(2.6)
sec.header_distance = sec.footer_distance = Cm(1.3)
sec.different_first_page_header_footer = True

def font(style, name, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)

font(doc.styles['Normal'], '宋体', 12)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.5
normal.first_line_indent = Pt(24)
normal.space_after = Pt(5)
normal.widow_control = True
for level, size in [(1,16),(2,14),(3,12)]:
    s = doc.styles[f'Heading {level}']
    font(s, '黑体', size)
    s.paragraph_format.first_line_indent = Pt(0)
    s.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.line_spacing = 1.25
    s.paragraph_format.keep_with_next = True
font(doc.styles['Title'], '黑体', 22)
doc.styles['Title'].paragraph_format.first_line_indent = Pt(0)
font(doc.styles['Subtitle'], '宋体', 16)
doc.styles['Subtitle'].font.italic = False
doc.styles['Subtitle'].paragraph_format.first_line_indent = Pt(0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hp.paragraph_format.first_line_indent = Pt(0)
hr = hp.add_run('智聘通招聘网站的设计与实现  开题报告')
hr.font.size = Pt(9)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Pt(0)
footer.add_run('— ')
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
footer.add_run(' —')

lines = source.read_text(encoding='utf-8').splitlines()
doc.add_paragraph('本科毕业设计开题报告', 'Subtitle').alignment = WD_ALIGN_PARAGRAPH.CENTER
title = doc.add_paragraph('智聘通招聘网站的设计与实现', 'Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(80)
title.paragraph_format.space_after = Pt(62)
for line in lines[4:next(k for k,l in enumerate(lines) if l.startswith('# 一'))]:
    if line.strip():
        p = doc.add_paragraph(line)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(18)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

def add_table(rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for c,w in zip(table.columns,[2.8,2.7,10.3]): c.width = Cm(w)
    for i,row in enumerate(rows):
        cells = table.rows[0].cells if i == 0 else table.add_row().cells
        for c,w,t in zip(cells,[2.8,2.7,10.3],row):
            c.width = Cm(w); c.text = t
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e = OxmlElement(f'w:{side}'); e.set(qn('w:w'),'100'); e.set(qn('w:type'),'dxa'); margins.append(e)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement(f'w:{side}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            tcpr.append(borders)
            if i==0:
                shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'E7E6E6'); tcpr.append(shading)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                for r in p.runs:
                    r.font.size=Pt(10.5); r.bold=(i==0)
                if c is cells[0] or len(t)<12: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        trpr=table.rows[i]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        if i==0: repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

i=next(k for k,l in enumerate(lines) if l.startswith('# 一'))
refs=False
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[- :]+',x) for x in row): rows.append(row)
            i+=1
        add_table(rows); continue
    m=re.match(r'^(#{1,3}) (.+)$',line)
    if m:
        if m.group(2).startswith('十 '):
            refs=True
            doc.add_page_break()
        p=doc.add_heading(m.group(2),level=len(m.group(1)))
    else:
        p=doc.add_paragraph(line)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if refs:
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_together=True
            p.paragraph_format.first_line_indent=Pt(-18)
            p.paragraph_format.left_indent=Pt(18)
            p.paragraph_format.line_spacing=1.2
            p.paragraph_format.space_after=Pt(7)
            for r in p.runs: r.font.size=Pt(10.5)
    i+=1
doc.core_properties.title='智聘通招聘网站的设计与实现'
doc.core_properties.subject='本科毕业设计开题报告'
doc.core_properties.author=''
out.parent.mkdir(parents=True,exist_ok=True)
doc.save(out)
(root/'references/prompt/background.md').write_text('''# 课题背景

课题名称：智聘通招聘网站的设计与实现。
任务范围：根据现有代码与项目文件撰写独立开题报告，不执行完整毕业论文分章创作和检测流程。
工程来源：G:/project/CareerNova。已有求职者、企业和管理员三角色门户，采用 Java 17、Spring Boot 3.4.5、MyBatis-Plus、Spring Security、MySQL 8、Vue 3、TypeScript、Vite、Pinia、Element Plus、ECharts。
业务：岗位检索、简历管理与文件解析、投递状态处理、收藏、举报与日志；智能辅助包括简历分析、岗位匹配、推荐、优化、模拟面试和助手。
数据库：schema.sql 含16张表，不设物理外键，Service维护逻辑关系。算法实际权重为语义50%、技能25%、经验15%、学历10%。本地向量为128维2-gram哈希归一化，不能等同于真实语义模型。
状态：项目记载2026-10-08历史测试结果，本次只核对源码与文件；真实模型、生产部署等尚无完整验收，需作为拟开展任务。学校、学院、作者、教师、日期未知，封面留空。进度按开题后12周拟定。
报告依据：README.md、docs/项目总文档.md、docs/系统功能说明.md、docs/人工智能功能说明.md、docs/数据库设计.md、docs/升级说明与验收.md，及后端匹配、推荐、投递、安全与文件相关源码。
''',encoding='utf-8')
shutil.copy2(out,Path(r'G:\project')/out.name)
text=source.read_text(encoding='utf-8')
print(json.dumps({'output':str(Path(r'G:\project')/out.name),'characters':len(text),'paragraphs':len(doc.paragraphs),'references':len(re.findall(r'^\[\d+\]',text,re.M))},ensure_ascii=False))
